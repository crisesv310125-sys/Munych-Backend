"""Exportación a Excel."""

import io
import uuid
from datetime import UTC, date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from munych_api.modules.results.application.ports import (
    CommitteeResults,
    ExportData,
    ScoreDetail,
)
from munych_api.modules.results.domain.models import (
    Award,
    CommitteeInfo,
    DelegateScore,
    DelegationInfo,
    RankingEntry,
)
from munych_api.modules.results.infrastructure.excel_exporter import ExcelExporter


def test_excel_has_summary_committee_and_detail_sheets() -> None:
    mx = DelegationInfo(uuid.uuid4(), "MX", None, "=Colegio malicioso", ())
    entry = RankingEntry(
        position=1,
        tied=False,
        delegation=mx,
        delegates=(DelegateScore(uuid.uuid4(), "Ana Pérez", Decimal("87.456")),),
        score=Decimal("87.456"),
        criteria_averages={},
    )
    committee = CommitteeInfo(
        uuid.uuid4(),
        uuid.uuid4(),
        "disec",
        "Primera Comisión",
        "DISEC/1",
        "general_assembly",
        {},
        True,
    )
    data = ExportData(
        event_name="SD-MUN 2026",
        venue="Santo Domingo",
        starts_on=date(2026, 10, 9),
        ends_on=date(2026, 10, 11),
        timezone="America/Santo_Domingo",
        generated_at=datetime(2026, 10, 11, 23, 0, tzinfo=UTC),
        committees=[CommitteeResults(committee, [entry], [Award("best_delegate", mx.id)])],
        details=[
            ScoreDetail(
                committee="Primera Comisión",
                delegation="MX",
                delegate="Ana Pérez",
                institution="Colegio",
                session_number=1,
                judge="Laura Méndez",
                criterion="Oratoria",
                value=4,
                comment="Muy bien",
                updated_at=datetime(2026, 10, 10, 15, 0, tzinfo=UTC),
            )
        ],
    )
    content = ExcelExporter().export(data)
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Resumen", "DISEC-1", "Detalle"]

    summary = workbook["Resumen"]
    assert summary["A1"].value == "SD-MUN 2026"
    assert [c.value for c in summary[7]] == [
        "Comité",
        "Premio",
        "Posición",
        "Delegación",
        "Delegados",
        "Institución",
        "Nota",
    ]
    assert [c.value for c in summary[8]] == [
        "Primera Comisión",
        "Mejor delegado",
        1,
        "MX",
        "Ana Pérez",
        "'=Colegio malicioso",  # sin inyección de fórmulas
        87.46,
    ]

    ranking = workbook["DISEC-1"]
    assert [c.value for c in ranking[4]][:7] == [
        1,
        None,
        "MX",
        "'=Colegio malicioso",
        "Ana Pérez",
        87.46,
        "Mejor delegado",
    ]

    detail = workbook["Detalle"]
    assert detail["J1"].value == "Actualizado (America/Santo_Domingo)"
    # 15:00 UTC = 11:00 en Santo Domingo (UTC−4).
    assert detail["J2"].value == datetime(2026, 10, 10, 11, 0)  # noqa: DTZ001 — Excel no guarda zona
    assert [c.value for c in detail[2]][:9] == [
        "Primera Comisión",
        "MX",
        "Ana Pérez",
        "Colegio",
        1,
        "Laura Méndez",
        "Oratoria",
        4,
        "Muy bien",
    ]
