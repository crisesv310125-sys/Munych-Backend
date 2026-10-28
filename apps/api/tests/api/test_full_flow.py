"""Flujo completo de la sección 10, de extremo a extremo contra la BD local.

Crear evento → comité → importar CSV → invitar y asignar juez → abrir sesión →
calificar → cerrar sesión → ranking → premios → cerrar evento → publicar →
la vista previa coincide con `public_results`.
"""

import io
import uuid
from datetime import UTC, datetime
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine

from tests.support.auth import bearer, make_token
from tests.support.seed import ADMIN_ID, JUDGE_2_ID
from tests.support.supabase import FakeRevalidator

pytestmark = pytest.mark.integration

API = "/api/v1"
ADMIN = bearer(make_token(ADMIN_ID, role="admin"))
OTHER_JUDGE = bearer(make_token(JUDGE_2_ID, role="judge"))


def _ok(response: Any, status: int = 200) -> Any:
    assert response.status_code == status, response.text
    return response.json()


async def test_full_flow(
    client: AsyncClient, engine: AsyncEngine, fake_revalidator: FakeRevalidator
) -> None:
    slug = f"flujo-{uuid.uuid4().hex[:8]}"

    # 1. Evento y comité (1 sesión, delegación individual).
    event = _ok(
        await client.post(
            f"{API}/events",
            headers=ADMIN,
            json={
                "slug": slug,
                "name": "Flujo MUN",
                "brand": "munych",
                "venue": "CDMX",
                "timezone": "America/Mexico_City",
                "starts_on": "2026-11-01",
                "ends_on": "2026-11-02",
            },
        ),
        201,
    )
    committee = _ok(
        await client.post(
            f"{API}/events/{event['id']}/committees",
            headers=ADMIN,
            json={
                "name": "Primera Comisión",
                "acronym": "DISEC",
                "slug": "disec",
                "type": "general_assembly",
                "delegation_mode": "single",
                "sessions_count": 1,
            },
        ),
        201,
    )
    committee_id = committee["id"]

    # 2. Importación CSV.
    csv_body = (
        "country_code,character_name,institution,delegate_1,delegate_2\n"
        "MX,,Colegio Ejemplo,Ana Pérez López,\n"
        "US,,Instituto Norte,John Smith,\n"
    )
    imported = _ok(
        await client.post(
            f"{API}/committees/{committee_id}/delegations/import",
            headers=ADMIN,
            files={"file": ("d.csv", csv_body.encode(), "text/csv")},
        )
    )
    assert imported == {"created": 2, "skipped": 0, "errors": []}

    # 3. Invitar juez y asignarlo.
    judge = _ok(
        await client.post(
            f"{API}/judges",
            headers=ADMIN,
            json={"email": f"{slug}@munych.dev", "full_name": "Jueza Flujo"},
        ),
        201,
    )
    judge_headers = bearer(make_token(judge["judge_id"], role="judge"))
    _ok(
        await client.put(
            f"{API}/committees/{committee_id}/judges",
            headers=ADMIN,
            json=[{"judge_id": judge["judge_id"], "role": "chair"}],
        )
    )

    # 4. El juez ve su comité solo cuando el evento está activo; la sesión empieza cerrada.
    assert _ok(await client.get(f"{API}/judge/committees", headers=judge_headers)) == []
    _ok(await client.patch(f"{API}/events/{event['id']}", headers=ADMIN, json={"status": "active"}))
    assigned = _ok(await client.get(f"{API}/judge/committees", headers=judge_headers))
    assert [c["committee_id"] for c in assigned] == [committee_id]
    assert assigned[0]["event"]["name"] == "Flujo MUN"
    detail = _ok(await client.get(f"{API}/judge/committees/{committee_id}", headers=judge_headers))
    session_id = detail["sessions"][0]["session_id"]
    criteria = detail["rubric"]["criteria"]
    delegates = {d["country_code"]: d["delegates"][0]["delegate_id"] for d in detail["delegations"]}

    def batch(values: dict[str, int]) -> list[dict[str, Any]]:
        return [
            {
                "client_id": str(uuid.uuid4()),
                "delegate_id": delegates[country],
                "session_id": session_id,
                "criterion_id": c["id"],
                "value": value,
                "comment": "Comentario privado del juez",
                "client_updated_at": datetime.now(UTC).isoformat(),
            }
            for country, value in values.items()
            for c in criteria
        ]

    closed = _ok(
        await client.post(
            f"{API}/judge/scores/batch", headers=judge_headers, json={"items": batch({"MX": 4})}
        )
    )
    assert {r["code"] for r in closed["rejected"]} == {"SESSION_CLOSED"}

    # 5. Abrir sesión y calificar (con un ítem fuera de rango y un juez no asignado).
    _ok(await client.post(f"{API}/committees/{committee_id}/sessions/1/open", headers=ADMIN))
    items = batch({"MX": 4, "US": 3})
    bad = {**items[0], "client_id": str(uuid.uuid4()), "value": 9}
    result = _ok(
        await client.post(
            f"{API}/judge/scores/batch", headers=judge_headers, json={"items": [*items, bad]}
        )
    )
    assert len(result["accepted"]) == len(items)
    assert [(r["client_id"], r["code"]) for r in result["rejected"]] == [
        (bad["client_id"], "VALUE_OUT_OF_RANGE")
    ]
    again = _ok(
        await client.post(f"{API}/judge/scores/batch", headers=judge_headers, json={"items": items})
    )
    assert len(again["accepted"]) == len(items)  # idempotente

    intruder = _ok(
        await client.post(
            f"{API}/judge/scores/batch", headers=OTHER_JUDGE, json={"items": items[:1]}
        )
    )
    assert intruder["rejected"][0]["code"] == "JUDGE_NOT_ASSIGNED"
    forbidden = await client.get(f"{API}/judge/sessions/{session_id}/scores", headers=OTHER_JUDGE)
    assert forbidden.status_code == 403

    mine = _ok(await client.get(f"{API}/judge/sessions/{session_id}/scores", headers=judge_headers))
    assert len(mine) == len(items)

    progress = _ok(await client.get(f"{API}/committees/{committee_id}/progress", headers=ADMIN))
    assert progress["sessions"][0]["judges"] == [
        {
            "judge_id": judge["judge_id"],
            "full_name": "Jueza Flujo",
            "scored": len(items),
            "expected": 2 * len(criteria),
        }
    ]

    # 6. No se puede cerrar el evento con sesiones abiertas.
    still_open = await client.post(f"{API}/events/{event['id']}/close", headers=ADMIN)
    assert still_open.status_code == 409
    assert still_open.json()["code"] == "SESSIONS_STILL_OPEN"
    _ok(await client.post(f"{API}/committees/{committee_id}/sessions/1/close", headers=ADMIN))

    # 7. Ranking y premios (con aviso de cupo).
    ranking = _ok(await client.get(f"{API}/committees/{committee_id}/ranking", headers=ADMIN))
    assert [(e["country_code"], e["position"], e["score"]) for e in ranking["items"]] == [
        ("MX", 1, 100.0),
        ("US", 2, 75.0),
    ]
    mx_id, us_id = (e["delegation_id"] for e in ranking["items"])
    awards = _ok(
        await client.put(
            f"{API}/committees/{committee_id}/awards",
            headers=ADMIN,
            json=[
                {"award_type": "best_delegate", "delegation_id": mx_id},
                {"award_type": "best_delegate", "delegation_id": us_id, "position": 2},
            ],
        )
    )
    assert awards["quota_warnings"] == ["Mejor delegado: 2 asignados y el cupo es 1"]
    _ok(
        await client.put(
            f"{API}/committees/{committee_id}/awards",
            headers=ADMIN,
            json=[
                {"award_type": "best_delegate", "delegation_id": mx_id},
                {"award_type": "outstanding_delegate", "delegation_id": us_id},
            ],
        )
    )

    # 8. Publicar antes de cerrar falla; cerrar crea la instantánea.
    early = await client.post(f"{API}/events/{event['id']}/publish", headers=ADMIN)
    assert early.json()["code"] == "EVENT_NOT_CLOSED"
    closed_event = _ok(await client.post(f"{API}/events/{event['id']}/close", headers=ADMIN))
    assert closed_event["status"] == "closed"
    assert _ok(await client.get(f"{API}/judge/committees", headers=judge_headers)) == []

    late = _ok(
        await client.post(
            f"{API}/judge/scores/batch", headers=judge_headers, json={"items": batch({"MX": 1})}
        )
    )
    assert {r["code"] for r in late["rejected"]} == {"EVENT_CLOSED"}
    frozen = await client.put(f"{API}/committees/{committee_id}/awards", headers=ADMIN, json=[])
    assert frozen.json()["code"] == "EVENT_CLOSED"

    # 9. Excel.
    export = await client.get(f"{API}/events/{event['id']}/export.xlsx", headers=ADMIN)
    assert export.status_code == 200
    assert f"{slug}-resultados.xlsx" in export.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(export.content))
    assert workbook.sheetnames == ["Resumen", "DISEC", "Detalle"]
    assert workbook["Detalle"].max_row == 1 + len(items)

    # 10. Vista previa → publicar → coincide con public_results.
    preview = _ok(
        await client.get(f"{API}/events/{event['id']}/publication/preview", headers=ADMIN)
    )
    assert preview["committees"][0]["awards"][0]["delegates"] == ["Ana P."]
    assert preview["committees"][0]["ranking"] is None
    published = _ok(await client.post(f"{API}/events/{event['id']}/publish", headers=ADMIN))
    assert (published["version"], published["revalidated"]) == (1, True)
    assert fake_revalidator.calls == [slug]

    async with engine.connect() as conn:
        row = (
            await conn.execute(
                text(
                    "select published, version, data from public.public_results"
                    " where event_slug = :s"
                ),
                {"s": slug},
            )
        ).one()
    assert row.published is True
    assert row.version == 1
    stored = dict(row.data)
    assert stored.pop("published_at")
    preview.pop("published_at")
    assert stored == preview
    text_data = str(row.data)
    assert "Comentario privado" not in text_data
    assert "Pérez" not in text_data
    assert judge["judge_id"] not in text_data

    # 11. Cambiar la configuración y volver a publicar sube la versión; despublicar avisa.
    _ok(
        await client.put(
            f"{API}/events/{event['id']}/publication-settings",
            headers=ADMIN,
            json={"name_format": "full", "show_ranking": True, "show_scores": True},
        )
    )
    republished = _ok(await client.post(f"{API}/events/{event['id']}/publish", headers=ADMIN))
    assert republished["version"] == 2
    unpublished = _ok(await client.post(f"{API}/events/{event['id']}/unpublish", headers=ADMIN))
    assert unpublished == {"revalidated": True}

    async with engine.connect() as conn:
        audit_count = await conn.scalar(
            text("select count(*) from app.audit_log where actor_id = :j and entity = 'score'"),
            {"j": judge["judge_id"]},
        )
    assert audit_count == len(items)
