"""Exportación a Excel (sección 8.3): Resumen, una hoja por comité y Detalle."""

import io
import re
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from munych_api.modules.results.application.ports import CommitteeResults, ExportData
from munych_api.modules.results.domain.awards import sort_awards
from munych_api.modules.results.domain.models import RankingEntry
from munych_api.scoring_engine import present
from munych_api.shared.kernel.catalog import AWARD_LABELS

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

Cell = str | int | float | datetime | None

_FLAG_LABELS = {
    "incomplete": "juez incompleto",
    "session_without_scores": "sesión sin notas",
    "delegate_without_scores": "delegado sin notas",
}


def _safe(value: Cell) -> Cell:
    """Evita que un texto importado (p. ej. del CSV) se interprete como fórmula."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _score(value: object) -> float | None:
    rounded = present(value)  # type: ignore[arg-type]
    return None if rounded is None else float(rounded)


def _local(value: datetime, timezone: str) -> datetime:
    # Excel no admite zonas horarias: se presenta en la hora local del evento.
    return value.astimezone(ZoneInfo(timezone)).replace(tzinfo=None)


def _write_table(
    sheet: Worksheet, start_row: int, headers: list[str], rows: list[list[Cell]]
) -> None:
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(row, start=1):
            sheet.cell(row=r, column=col, value=_safe(value))


def _autosize(sheet: Worksheet) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 10), 60)


def _sheet_title(name: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub("-", name).strip("'") or "Comité"
    base = base[:31]
    title, counter = base, 2
    while title.lower() in used:
        suffix = f" ({counter})"
        title = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(title.lower())
    return title


def _delegation_label(entry: RankingEntry) -> str:
    return entry.delegation.label


def _delegates(entry: RankingEntry) -> str:
    return ", ".join(d.full_name for d in entry.delegates)


class ExcelExporter:
    media_type = XLSX_MEDIA_TYPE
    extension = "xlsx"

    def export(self, data: ExportData) -> bytes:
        workbook = Workbook()
        summary = workbook.active
        assert summary is not None  # noqa: S101 — un libro nuevo siempre tiene hoja activa
        summary.title = "Resumen"
        used = {"resumen", "detalle"}
        self._summary(summary, data)
        for committee in data.committees:
            sheet = workbook.create_sheet(_sheet_title(committee.committee.acronym, used))
            self._ranking(sheet, committee)
        self._details(workbook.create_sheet("Detalle"), data)

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _summary(sheet: Worksheet, data: ExportData) -> None:
        sheet["A1"] = _safe(data.event_name)
        sheet["A1"].font = Font(bold=True, size=14)
        sheet["A2"] = "Sede"
        sheet["B2"] = _safe(data.venue)
        sheet["A3"] = "Fechas"
        sheet["B3"] = f"{data.starts_on.isoformat()} a {data.ends_on.isoformat()}"
        sheet["A4"] = "Zona horaria"
        sheet["B4"] = data.timezone
        sheet["A5"] = "Generado"
        sheet["B5"] = _local(data.generated_at, data.timezone)
        sheet["B5"].number_format = "yyyy-mm-dd hh:mm"

        rows: list[list[Cell]] = []
        for committee in data.committees:
            entries = {e.delegation.id: e for e in committee.ranking}
            for award in sort_awards(committee.awards):
                entry = entries.get(award.delegation_id)
                if entry is None:
                    continue
                rows.append(
                    [
                        committee.committee.name,
                        AWARD_LABELS[award.award_type],
                        award.position,
                        _delegation_label(entry),
                        _delegates(entry),
                        entry.delegation.institution,
                        _score(entry.score),
                    ]
                )
        _write_table(
            sheet,
            7,
            ["Comité", "Premio", "Posición", "Delegación", "Delegados", "Institución", "Nota"],
            rows,
        )
        _autosize(sheet)

    @staticmethod
    def _ranking(sheet: Worksheet, committee: CommitteeResults) -> None:
        info = committee.committee
        sheet["A1"] = _safe(f"{info.name} ({info.acronym})")
        sheet["A1"].font = Font(bold=True, size=13)
        awards_by_delegation: dict[object, list[str]] = {}
        for award in sort_awards(committee.awards):
            awards_by_delegation.setdefault(award.delegation_id, []).append(
                AWARD_LABELS[award.award_type]
            )
        rows: list[list[Cell]] = [
            [
                entry.position,
                "Sí" if entry.tied else "",
                _delegation_label(entry),
                entry.delegation.institution,
                _delegates(entry),
                _score(entry.score),
                ", ".join(awards_by_delegation.get(entry.delegation.id, [])),
                ", ".join(
                    sorted({_FLAG_LABELS.get(f.kind, f.kind) for f in entry.incomplete_flags})
                ),
            ]
            for entry in committee.ranking
        ]
        _write_table(
            sheet,
            3,
            [
                "Posición",
                "Empate",
                "Delegación",
                "Institución",
                "Delegados",
                "Nota",
                "Premios",
                "Avisos",
            ],
            rows,
        )
        _autosize(sheet)

    @staticmethod
    def _details(sheet: Worksheet, data: ExportData) -> None:
        rows: list[list[Cell]] = [
            [
                d.committee,
                d.delegation,
                d.delegate,
                d.institution,
                d.session_number,
                d.judge,
                d.criterion,
                d.value,
                d.comment,
                _local(d.updated_at, data.timezone),
            ]
            for d in data.details
        ]
        _write_table(
            sheet,
            1,
            [
                "Comité",
                "Delegación",
                "Delegado",
                "Institución",
                "Sesión",
                "Juez",
                "Criterio",
                "Valor",
                "Comentario",
                f"Actualizado ({data.timezone})",
            ],
            rows,
        )
        for (cell,) in sheet.iter_rows(min_row=2, min_col=10, max_col=10):
            cell.number_format = "yyyy-mm-dd hh:mm"
        sheet.freeze_panes = "A2"
        _autosize(sheet)
